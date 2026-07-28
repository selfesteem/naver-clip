#!/usr/bin/env python3
"""
네이버 검색광고 API 키워드 검색량 조회

사용법:
    python search_volume.py keywords.xlsx
    python search_volume.py keywords.xlsx --output volume_result.xlsx

인증 정보는 환경변수 또는 인자로 전달:
    export NAVER_API_KEY=...
    export NAVER_SECRET_KEY=...
    export NAVER_CUSTOMER_ID=...
"""

import os
import sys
import time
import hmac
import hashlib
import base64
import argparse
from pathlib import Path
from datetime import date

import requests
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from urllib.parse import quote_plus

API_BASE = "https://api.searchad.naver.com"
KEYWORD_TOOL_PATH = "/keywordstool"
BATCH_SIZE = 5
REQUEST_DELAY = 0.12
PROGRESS_EVERY = 100


# ── 인증 ──────────────────────────────────────────────────────────────────────

def _sign(secret_key: str, timestamp: str, method: str, path: str) -> str:
    message = f"{timestamp}.{method}.{path}"
    sig = hmac.new(secret_key.encode(), message.encode(), hashlib.sha256).digest()
    return base64.b64encode(sig).decode()


# ── API 호출 ──────────────────────────────────────────────────────────────────

def fetch_batch(batch: list[str], api_key: str, secret_key: str, customer_id: str) -> list[dict]:
    timestamp = str(int(time.time() * 1000))
    headers = {
        "X-Timestamp": timestamp,
        "X-API-KEY": api_key,
        "X-Customer": customer_id,
        "X-Signature": _sign(secret_key, timestamp, "GET", KEYWORD_TOOL_PATH),
    }
    # Naver 검색광고 API는 공백 없는 키워드만 허용 (이혼 소송 → 이혼소송)
    hint = ",".join(quote_plus(kw.replace(" ", "")) for kw in batch)
    url = f"{API_BASE}{KEYWORD_TOOL_PATH}?hintKeywords={hint}&showDetail=1"
    r = requests.get(url, headers=headers, timeout=15)
    if not r.ok:
        raise requests.HTTPError(
            f"{r.status_code} {r.reason} — 응답: {r.text[:300]}",
            response=r,
        )
    return r.json().get("keywordList", [])


def get_search_volumes(
    keywords: list[str], api_key: str, secret_key: str, customer_id: str
) -> tuple[dict[str, dict], dict[str, dict]]:
    """
    Returns:
        input_map  : {원본키워드 → api_row}  (입력 키워드)
        related_map: {키워드(공백없음) → api_row}  (추천 키워드 — 입력 목록 외)
    """
    input_nsp_set = {kw.replace(" ", "") for kw in keywords}
    input_map: dict[str, dict] = {}
    related_map: dict[str, dict] = {}
    total_batches = (len(keywords) + BATCH_SIZE - 1) // BATCH_SIZE

    for batch_idx, i in enumerate(range(0, len(keywords), BATCH_SIZE)):
        batch = keywords[i : i + BATCH_SIZE]
        # 공백제거 → 원본 역매핑
        nsp_to_orig = {kw.replace(" ", ""): kw for kw in batch}

        try:
            rows = fetch_batch(batch, api_key, secret_key, customer_id)
            for row in rows:
                rel = row.get("relKeyword", "")
                if not rel:
                    continue
                if rel in nsp_to_orig:
                    input_map[nsp_to_orig[rel]] = row
                elif rel not in input_nsp_set:
                    related_map.setdefault(rel, row)
        except Exception as e:
            print(f"  오류 (배치 {batch_idx + 1}/{total_batches}): {e}")
            for kw in batch:
                input_map.setdefault(kw, {"_error": str(e)})

        if (batch_idx + 1) % PROGRESS_EVERY == 0:
            print(f"  {i + len(batch)}/{len(keywords)} 처리...")

        if i + BATCH_SIZE < len(keywords):
            time.sleep(REQUEST_DELAY)

    return input_map, related_map


# ── 키워드 로드 ────────────────────────────────────────────────────────────────

def load_keywords(filepath: str, col: str | None) -> list[str]:
    path = Path(filepath)
    if not path.exists():
        sys.exit(f"파일을 찾을 수 없습니다: {filepath}")
    xl = pd.ExcelFile(path)
    df = xl.parse(xl.sheet_names[0], dtype=str)
    if df.empty:
        sys.exit("엑셀 파일이 비어있습니다.")
    if col and col in df.columns:
        kw_col = col
    elif "키워드" in df.columns:
        kw_col = "키워드"
    else:
        kw_col = df.columns[0]
    return [k for k in df[kw_col].dropna().str.strip().tolist() if k]


# ── 데이터 변환 ───────────────────────────────────────────────────────────────

def _safe_int(val) -> int | None:
    if val is None:
        return None
    s = str(val).strip()
    if s.startswith("<"):
        return 9  # API 명세: 10 미만이면 "<10" 반환
    try:
        return int(s)
    except ValueError:
        return None


def _row_to_record(display_kw: str, api_row: dict) -> dict:
    pc  = _safe_int(api_row.get("monthlyPcQcCnt"))
    mob = _safe_int(api_row.get("monthlyMobileQcCnt"))
    total = (pc or 0) + (mob or 0) if (pc is not None or mob is not None) else None
    return {
        "키워드":          display_kw,
        "월간PC검색수":     pc,
        "월간모바일검색수": mob,
        "월간검색수합계":   total,
        "경쟁정도":         api_row.get("compIdx", ""),
        "오류":             api_row.get("_error", ""),
    }


def build_dataframes(
    keywords: list[str],
    input_map: dict[str, dict],
    related_map: dict[str, dict],
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    # 1) 조회 키워드
    input_rows = [_row_to_record(kw, input_map.get(kw, {})) for kw in keywords]
    df_input = pd.DataFrame(input_rows)

    # 2) 추천 키워드 (검색량 내림차순)
    related_rows = [_row_to_record(kw, row) for kw, row in related_map.items()]
    df_related = pd.DataFrame(related_rows) if related_rows else pd.DataFrame(
        columns=["키워드", "월간PC검색수", "월간모바일검색수", "월간검색수합계", "경쟁정도", "오류"]
    )
    if not df_related.empty and "월간검색수합계" in df_related.columns:
        df_related = df_related.sort_values("월간검색수합계", ascending=False, na_position="last")

    # 3) 통계
    valid = df_input[df_input["월간검색수합계"].notna()].copy()
    stats_rows = [
        ("전체 조회 키워드",     len(df_input)),
        ("조회 성공",            len(valid)),
        ("조회 오류",            int(df_input["오류"].astype(bool).sum())),
        ("추천 키워드 수",       len(df_related)),
        ("", ""),
        ("검색량 평균 (합계)",   int(valid["월간검색수합계"].mean()) if not valid.empty else 0),
        ("검색량 최대",          int(valid["월간검색수합계"].max()) if not valid.empty else 0),
        ("검색량 최소",          int(valid["월간검색수합계"].min()) if not valid.empty else 0),
        ("", ""),
        ("경쟁 낮음",      int((valid["경쟁정도"] == "낮음").sum())),
        ("경쟁 중간",      int((valid["경쟁정도"] == "중간").sum())),
        ("경쟁 높음",      int((valid["경쟁정도"] == "높음").sum())),
        ("", ""),
        ("★ 검색량 1000+ & 경쟁 낮음", int(((valid["월간검색수합계"] >= 1000) & (valid["경쟁정도"] == "낮음")).sum())),
    ]
    df_stats = pd.DataFrame(stats_rows, columns=["항목", "값"])

    return df_input, df_related, df_stats


# ── 엑셀 저장 ─────────────────────────────────────────────────────────────────

HEADER_FILL  = PatternFill("solid", fgColor="2F5496")
SECTION_FILL = PatternFill("solid", fgColor="D6E4F7")
GOLD_FILL    = PatternFill("solid", fgColor="FFF2CC")
WHITE_FONT   = Font(bold=True, color="FFFFFF")
BOLD         = Font(bold=True)


def _style_header(ws, row: int, cols: int):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center")


def _auto_width(ws):
    for col_cells in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col_cells), default=8)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 40)


def _write_df(ws, df: pd.DataFrame, start_row: int, highlight_fn=None):
    for ci, col_name in enumerate(df.columns, 1):
        ws.cell(row=start_row, column=ci, value=col_name)
    _style_header(ws, start_row, len(df.columns))
    for ri, record in enumerate(df.itertuples(index=False), start_row + 1):
        for ci, val in enumerate(record, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            if highlight_fn and highlight_fn(record):
                cell.fill = GOLD_FILL


def save_excel(
    df_input: pd.DataFrame,
    df_related: pd.DataFrame,
    df_stats: pd.DataFrame,
    output_path: Path,
):
    def is_prime(rec):
        """검색량 1000+ & 경쟁 낮음 → 황색 강조"""
        total = getattr(rec, "월간검색수합계", None)
        comp  = getattr(rec, "경쟁정도", "")
        return total is not None and total >= 1000 and comp == "낮음"

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # ── 시트 1: 조회 키워드 ──
        df_input.to_excel(writer, index=False, sheet_name="조회키워드")
        ws1 = writer.sheets["조회키워드"]
        _style_header(ws1, 1, len(df_input.columns))
        for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
            rec_vals = [c.value for c in row]
            total = rec_vals[3]  # 월간검색수합계
            comp  = rec_vals[4]  # 경쟁정도
            if total is not None and isinstance(total, (int, float)) and total >= 1000 and comp == "낮음":
                for c in row:
                    c.fill = GOLD_FILL
        _auto_width(ws1)

        # ── 시트 2: 추천 키워드 ──
        if not df_related.empty:
            df_related.to_excel(writer, index=False, sheet_name="추천키워드")
            ws2 = writer.sheets["추천키워드"]
            _style_header(ws2, 1, len(df_related.columns))
            for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
                rec_vals = [c.value for c in row]
                total = rec_vals[3]
                comp  = rec_vals[4]
                if total is not None and isinstance(total, (int, float)) and total >= 1000 and comp == "낮음":
                    for c in row:
                        c.fill = GOLD_FILL
            _auto_width(ws2)

        # ── 시트 3: 통계 ──
        df_stats.to_excel(writer, index=False, sheet_name="통계")
        ws3 = writer.sheets["통계"]
        _style_header(ws3, 1, 2)
        for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
            label = row[0].value or ""
            if "★" in str(label):
                for c in row:
                    c.fill = GOLD_FILL
                    c.font = BOLD
        _auto_width(ws3)

    print(f"저장: {output_path}  (시트: 조회키워드 / 추천키워드 / 통계)")


# ── 터미널 요약 ───────────────────────────────────────────────────────────────

def print_summary(df_input: pd.DataFrame, df_related: pd.DataFrame):
    valid = df_input[df_input["월간검색수합계"].notna()]
    total = len(df_input)
    print(f"\n--- 요약 ({len(valid)}/{total} 성공) ---")
    print(f"추천 키워드: {len(df_related)}개")

    prime = valid[(valid["월간검색수합계"] >= 1000) & (valid["경쟁정도"] == "낮음")]
    if not prime.empty:
        print(f"\n[★ 검색량 1000+ & 경쟁 낮음 — {len(prime)}개]")
        print(prime.nlargest(10, "월간검색수합계")[["키워드", "월간검색수합계"]].to_string(index=False))

    if not valid.empty:
        print(f"\n[검색량 상위 5개]")
        print(valid.nlargest(5, "월간검색수합계")[["키워드", "월간검색수합계", "경쟁정도"]].to_string(index=False))


# ── 진입점 ────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="네이버 검색광고 API 키워드 검색량 조회")
    parser.add_argument("input_file", help="키워드 Excel 파일 경로")
    parser.add_argument("--col",         help="키워드 컬럼명 (기본: '키워드' 또는 첫 번째 컬럼)")
    parser.add_argument("--output",      help="결과 파일 경로 (기본: 입력 파일 폴더에 자동 생성)")
    parser.add_argument("--api-key",     default=os.environ.get("NAVER_API_KEY"))
    parser.add_argument("--secret-key",  default=os.environ.get("NAVER_SECRET_KEY"))
    parser.add_argument("--customer-id", default=os.environ.get("NAVER_CUSTOMER_ID"))
    parser.add_argument("--check-auth",  action="store_true", help="인증 정보 확인 후 종료")
    args = parser.parse_args()

    missing = [name for name, val in [
        ("NAVER_API_KEY",     args.api_key),
        ("NAVER_SECRET_KEY",  args.secret_key),
        ("NAVER_CUSTOMER_ID", args.customer_id),
    ] if not val]
    if missing:
        sys.exit(
            f"인증 정보 누락: {', '.join(missing)}\n"
            "환경변수(NAVER_API_KEY / NAVER_SECRET_KEY / NAVER_CUSTOMER_ID) 또는\n"
            "--api-key / --secret-key / --customer-id 인자로 전달하세요.\n\n"
            "발급: https://searchad.naver.com → 도구 → API 사용 관리"
        )

    if args.check_auth:
        print("[인증 정보 확인]")
        print(f"  NAVER_API_KEY     : {args.api_key[:8]}...{args.api_key[-4:]}  (길이 {len(args.api_key)})")
        print(f"  NAVER_SECRET_KEY  : {args.secret_key[:4]}...{args.secret_key[-4:]}  (길이 {len(args.secret_key)})")
        print(f"  NAVER_CUSTOMER_ID : {args.customer_id}")
        sys.exit(0)

    keywords = load_keywords(args.input_file, args.col)
    if not keywords:
        sys.exit("키워드가 없습니다.")

    today = date.today().strftime("%Y%m%d")
    input_path = Path(args.input_file)
    output_path = Path(args.output) if args.output else \
        input_path.parent / f"search_volume_{input_path.stem}_{today}.xlsx"

    eta_min = int(len(keywords) / BATCH_SIZE * REQUEST_DELAY // 60)
    print(f"키워드 {len(keywords)}개 조회 시작 (예상 소요: 약 {eta_min}분)...")

    input_map, related_map = get_search_volumes(
        keywords, args.api_key, args.secret_key, args.customer_id
    )

    df_input, df_related, df_stats = build_dataframes(keywords, input_map, related_map)
    save_excel(df_input, df_related, df_stats, output_path)
    print_summary(df_input, df_related)


if __name__ == "__main__":
    main()
